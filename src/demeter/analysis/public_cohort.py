"""Audit a pinned public cohort without estimating or activating health hazards."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from demeter.data.ingest import digest
from demeter.data.nhanes import encoded, storage_numbers
from demeter.schema import EvidenceRegistry

DATASET = "chen_public_intake"


def read_workbook(path: Path, spec: dict) -> pd.DataFrame:
    """Read values without recalculation; reject formulas rather than use caches."""
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if spec["sheet"] not in book.sheetnames:
            raise ValueError("Missing protocol sheet")
        rows = book[spec["sheet"]].iter_rows()
        headers = [c.value for c in next(rows)]
        names = spec["columns"]
        if any(headers.count(v) != 1 for v in names.values()):
            raise ValueError("Missing or duplicate protocol column")
        indices = [headers.index(v) for v in names.values()]
        values = []
        for row in rows:
            if any(c.data_type == "f" for c in row):
                raise ValueError("Workbook formulas require a separately reviewed extraction")
            values.append([row[i].value for i in indices])
        return pd.DataFrame(values, columns=list(names))
    finally:
        book.close()


def analyze_cohort(frame: pd.DataFrame, spec: dict) -> dict:
    """Report archive coverage and endpoint agreement, with unknowns retained."""
    if set(frame.columns) != set(spec["columns"]):
        raise ValueError("Unexpected cohort fields")
    if len(frame) != spec["published_records"]:
        raise ValueError("Cohort row count differs from the pinned protocol")
    data = frame.copy()
    missing = {k: int(data[k].isna().sum()) for k in data}
    invalid = {}
    for key in data.columns.drop("id"):
        numeric = pd.to_numeric(data[key], errors="coerce")
        bad = data[key].notna() & (~np.isfinite(numeric) | numeric.isna())
        if key in ("baseline_fpg", "final_fpg", "followup"):
            bad |= numeric.notna() & (numeric <= 0)
        elif key == "age":
            bad |= numeric.notna() & (numeric < spec["minimum_baseline_age_years"])
        elif key == "sex":
            bad |= numeric.notna() & ~numeric.isin(spec["sex_codes"].values())
        else:
            bad |= numeric.notna() & ~numeric.isin(spec["binary_codes"])
        invalid[key] = int(bad.sum())
        data[key] = numeric.mask(bad)
    blank_ids = data["id"].astype(str).str.strip().eq("")
    missing_ids = data["id"].isna() | blank_ids
    duplicate_rows = int(data.loc[~missing_ids, "id"].duplicated(keep=False).sum())
    threshold = spec["glucose_threshold_mmol_l"]
    positive = data["final_fpg"].ge(threshold) | data["diagnosis"].eq(1)
    negative = data["final_fpg"].lt(threshold) & data["diagnosis"].eq(0)
    known = positive | negative
    endpoint_known = data["endpoint"].notna()
    strict_disagreement = known & endpoint_known & (positive != data["endpoint"].eq(1))
    # A blank diagnosis is not a documented negative. Show this assumption separately.
    assumed_negative = data["diagnosis"].eq(0) | frame["diagnosis"].isna()
    conditional_known = positive | (data["final_fpg"].lt(threshold) & assumed_negative)
    conditional_disagreement = (
        conditional_known & endpoint_known & (positive != data["endpoint"].eq(1))
    )

    def distribution(series: pd.Series) -> dict:
        valid = series.dropna()
        return {
            "observed": len(valid),
            "missing_or_invalid": int(series.isna().sum()),
            "quantiles_years": [
                float(valid.quantile(q, interpolation=spec["quantile_method"]))
                if len(valid)
                else None
                for q in spec["quantiles"]
            ],
        }

    followup = {"all": distribution(data["followup"])}
    for code in spec["binary_codes"]:
        followup[f"recorded_endpoint_{code}"] = distribution(
            data.loc[data["endpoint"].eq(code), "followup"]
        )
    sex_counts = {k: int(data["sex"].eq(v).sum()) for k, v in spec["sex_codes"].items()}
    median = data["followup"].median()
    checks = {
        "published_sex_counts": all(sex_counts[k] == spec[f"published_{k}"] for k in sex_counts),
        "published_event_count": int(data["endpoint"].eq(1).sum())
        == spec["published_diabetes_events"],
        "published_median_followup": bool(np.isfinite(median))
        and round(float(median), spec["published_followup_decimal_places"])
        == spec["published_median_followup_years"],
        "valid_codes_and_values": not any(invalid.values()),
        "unique_present_identifiers": duplicate_rows == 0 and not missing_ids.any(),
        "required_fields_present": all(
            missing[k] == 0 for k in ["age", "sex", "baseline_fpg", "endpoint", "followup"]
        ),
        "baseline_diabetes_glucose_excluded": not data["baseline_fpg"].ge(threshold).any(),
        "known_endpoint_rule_agrees": not strict_disagreement.any(),
    }
    return storage_numbers(
        {
            "records": len(data),
            "sex_counts": sex_counts,
            "missing": missing,
            "invalid": invalid,
            "missing_or_blank_identifiers": int(missing_ids.sum()),
            "rows_with_duplicated_identifiers": duplicate_rows,
            "diagnosis_field": {
                "zero": int(data["diagnosis"].eq(0).sum()),
                "one": int(data["diagnosis"].eq(1).sum()),
                "missing": missing["diagnosis"],
                "invalid": invalid["diagnosis"],
            },
            "recorded_endpoint": {
                str(v): int(data["endpoint"].eq(v).sum()) for v in spec["binary_codes"]
            },
            "endpoint_reconciliation": {
                "observed_positive_rule": int(positive.sum()),
                "observed_negative_rule": int(negative.sum()),
                "unknown_rule": int((~known).sum()),
                "known_rule_compared": int((known & endpoint_known).sum()),
                "known_rule_disagreements": int(strict_disagreement.sum()),
                "if_blank_diagnosis_meant_negative": {
                    "compared": int((conditional_known & endpoint_known).sum()),
                    "unknown": int((~conditional_known).sum()),
                    "disagreements": int(conditional_disagreement.sum()),
                    "interpretation": "Conditional reconstruction only; blank coding unconfirmed",
                },
            },
            "followup": followup,
            "checks": {k: bool(v) for k, v in checks.items()},
            "source_reproduction_passed": all(checks.values()),
            "sampling_uncertainty": "Not estimated: exact fixed-release audit, not population inference",
        }
    )


def audit_public_cohort(registry: EvidenceRegistry, workbook: Path) -> dict:
    spec = registry.datasets[DATASET]
    if spec["model_role"] != "benchmark_only":
        raise ValueError("Public cohort intake must remain benchmark_only")
    protocol_bytes = Path(spec["protocol_path"]).read_bytes()
    protocol = json.loads(protocol_bytes)
    source = registry.sources[spec["source_id"]]
    if (
        digest(protocol_bytes) != spec["protocol_sha256"]
        or spec["analysis"] != protocol["analysis"]
        or spec["source_id"] != protocol["source_id"]
        or source.sha256 != protocol["archive_sha256"]
    ):
        raise ValueError("Public cohort protocol/registry mismatch")
    if digest(workbook.read_bytes()) != source.sha256:
        raise ValueError("Public cohort workbook checksum mismatch")
    result = analyze_cohort(read_workbook(workbook, spec["analysis"]), spec["analysis"])
    return {
        "schema_version": 1,
        "analysis_id": protocol["analysis_id"],
        "model_role": "benchmark_only",
        "scientific_release_ready": False,
        "assessment": "Public source reproduction; no transition hazard or health effect estimated",
        "population": spec["population"],
        "geography": spec["geography"],
        "time_period": spec["time_period"],
        "results": result,
        "limitations": protocol["limits"],
        "provenance": {
            "source": source.model_dump(mode="json"),
            "protocol_sha256": digest(protocol_bytes),
            "definition_sha256": digest(encoded(spec)),
            "transform_sha256": digest(Path(__file__).read_bytes()),
            "individual_results_exported": False,
            "raw_redistributed": False,
        },
    }
